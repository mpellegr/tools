# MPM Conversor Excel a CSV para After Effects

import csv
import os
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from openpyxl import load_workbook

APP_NAME = "MediaMotionLab | Excel to CSV"
VERSION = "v01.02"
DESCRIPTION = "Conversor Excel a CSV (UTF-8) para After Effects"

def xlsx_to_csv_ae_friendly(xlsx_path: str, sheet_name: str, csv_path: str) -> None:
    """
    AE-friendly:
    - delimitador fijo: coma
    - encoding: UTF-8 (sin BOM)
    - columnas fijas por fila (usa max_column)
    """
    wb = load_workbook(filename=xlsx_path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise ValueError(f"La hoja '{sheet_name}' no existe en el archivo.")

    ws = wb[sheet_name]
    ncols = ws.max_column or 0
    if ncols <= 0:
        wb.close()
        raise ValueError("No se detectaron columnas. Revisa que el XLSX tenga datos.")

    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(
            f,
            delimiter=",",
            quotechar='"',
            quoting=csv.QUOTE_MINIMAL,
            lineterminator="\n",
        )

        for row in ws.iter_rows(values_only=True):
            vals = ["" if v is None else v for v in row]
            if len(vals) < ncols:
                vals.extend([""] * (ncols - len(vals)))
            else:
                vals = vals[:ncols]
            writer.writerow(vals)

    wb.close()


class App(tk.Tk):
    def __init__(self):
        super().__init__()

        self.title(f"{APP_NAME} {VERSION}")
        self.geometry("512x360")
        self.resizable(False, False)

        self.xlsx_path = tk.StringVar(value="")
        self.sheet_name = tk.StringVar(value="")
        self.sheets = []

        self._build_ui()
        self._reset_ui(initial=True)

    def _build_ui(self):
        root = ttk.Frame(self, padding=22)
        root.pack(fill="both", expand=True)

        # Título / descripción
        title = ttk.Label(root, text=DESCRIPTION, font=("Segoe UI", 12))
        title.pack(pady=(0, 6))

        subtitle = ttk.Label(root, text=f"{VERSION}", font=("Segoe UI", 10))
        subtitle.pack(pady=(0, 18))

        # Bloque selección archivo
        block = ttk.Frame(root)
        block.pack(fill="x")

        ttk.Label(block, text="Selecciona archivo de Excel (.xlsx):").grid(row=0, column=0, sticky="w", pady=(0, 6))

        self.entry = ttk.Entry(block, textvariable=self.xlsx_path, width=62)
        self.entry.grid(row=1, column=0, sticky="w")

        btn = ttk.Button(block, text="Elegir…", command=self.pick_file)
        btn.grid(row=1, column=1, padx=(10, 0))

        # Bloque selección hoja
        ttk.Label(block, text="Selecciona la hoja:").grid(row=2, column=0, sticky="w", pady=(16, 6))

        self.sheet_combo = ttk.Combobox(block, textvariable=self.sheet_name, state="readonly", width=59, values=[])
        self.sheet_combo.grid(row=3, column=0, sticky="w")

        # Reaccionar al cambio de hoja para habilitar Convertir
        self.sheet_combo.bind("<<ComboboxSelected>>", lambda _e: self._update_convert_state())

        # Botón convertir
        convert_frame = ttk.Frame(root)
        convert_frame.pack(fill="x", pady=(25, 0))

        self.convert_btn = ttk.Button(convert_frame, text="CONVERTIR", command=self.convert)
        self.convert_btn.pack()

        # Estado
        self.status = ttk.Label(root, text="Listo.", font=("Segoe UI", 9))
        self.status.pack(side="left", pady=(0, 0))

    def _reset_ui(self, initial: bool = False):
        """
        Resetea elementos dependientes del archivo/hoja.
        """
        self.sheets = []
        self.sheet_combo["values"] = []
        self.sheet_name.set("")
        self.convert_btn.state(["disabled"])
        if initial:
            self.status.config(text="Selecciona un archivo XLSX…")
        else:
            self.status.config(text="Listo. Selecciona un archivo XLSX…")

    def _update_convert_state(self):
        """
        Habilita Convertir solo si hay archivo válido y hoja seleccionada.
        """
        xlsx = self.xlsx_path.get().strip()
        sheet = self.sheet_name.get().strip()
        ok = bool(xlsx) and os.path.isfile(xlsx) and bool(sheet)

        if ok:
            self.convert_btn.state(["!disabled"])
        else:
            self.convert_btn.state(["disabled"])

    def pick_file(self):
        path = filedialog.askopenfilename(
            title="Selecciona un archivo XLSX",
            filetypes=[("Excel (*.xlsx)", "*.xlsx")],
        )
        if not path:
            return

        # Reset antes de cargar, para no dejar estados antiguos
        self._reset_ui()

        self.xlsx_path.set(path)

        try:
            wb = load_workbook(filename=path, read_only=True, data_only=True)
            self.sheets = wb.sheetnames
            wb.close()
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo abrir el XLSX.\n\n{e}")
            self.xlsx_path.set("")
            self._reset_ui()
            return

        self.sheet_combo["values"] = self.sheets
        if self.sheets:
            self.sheet_name.set(self.sheets[0])
        else:
            self.sheet_name.set("")

        self.status.config(text="Archivo cargado. Selecciona la hoja y convierte.")
        self._update_convert_state()

    def convert(self):
        xlsx = self.xlsx_path.get().strip()
        sheet = self.sheet_name.get().strip()

        # Si el usuario llega aquí, debería estar OK, pero mantenemos guardas.
        if not xlsx or not os.path.isfile(xlsx):
            messagebox.showwarning("Falta archivo", "Selecciona un XLSX válido.")
            self._update_convert_state()
            return
        if not sheet:
            messagebox.showwarning("Falta hoja", "Selecciona una hoja.")
            self._update_convert_state()
            return

        default_name = os.path.splitext(os.path.basename(xlsx))[0] + f"_{sheet}.csv"
        csv_path = filedialog.asksaveasfilename(
            title="Guardar CSV",
            defaultextension=".csv",
            initialfile=default_name,
            filetypes=[("CSV (*.csv)", "*.csv")],
        )
        if not csv_path:
            return

        # Bloquear botón mientras convierte (por si el XLSX es grande)
        self.convert_btn.state(["disabled"])
        self.status.config(text="Convirtiendo…")

        try:
            xlsx_to_csv_ae_friendly(xlsx, sheet, csv_path)
        except Exception as e:
            messagebox.showerror("Error al convertir", f"No se pudo convertir.\n\n{e}")
            self.status.config(text="Conversión fallida.")
            self._update_convert_state()
            return

        self.status.config(text=f"OK: {os.path.basename(csv_path)}")
        messagebox.showinfo("Listo", "Conversión terminada.")
        self._update_convert_state()


if __name__ == "__main__":
    App().mainloop()
